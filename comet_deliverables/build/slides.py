from pathlib import Path
import json,sys,textwrap
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
B=Path(__file__).resolve().parent.parent
for name,file in [('DV','DejaVuSans.ttf'),('DVB','DejaVuSans-Bold.ttf')]:pdfmetrics.registerFont(TTFont(name,'/usr/share/fonts/truetype/dejavu/'+file))
W,H=13.333333,7.5
NAVY='11283F'; TEAL='087F8C'; GREEN='147D64'; MUTED='53667A'; LIGHT='EDF3F7'; GOLD='BC7412'; WHITE='FFFFFF'; RED='A53636'
slides=[]
def new(title,kicker='COMET-XAI LITE',source='',notes=''):
 s={'title':title,'elements':[],'notes':notes}
 slides.append(s)
 rect(s,0,0,W,H,WHITE);rect(s,0,0,.16,H,TEAL)
 text(s,.55,.27,12,.25,kicker,10,TEAL,True)
 text(s,.55,.73,12.2,.85,title,29,NAVY,True)
 rect(s,.55,6.96,12.2,.015,'CCD8E2')
 text(s,.55,7.08,11.4,.22,source or 'DTU • [Student name] • [Roll number] • Prof. Rahul Gupta • 28 October 2026',9,MUTED)
 text(s,12.1,7.06,.5,.25,str(len(slides)),10,MUTED)
 return s

def rect(s,x,y,w,h,c):s['elements'].append({'type':'rect','x':x,'y':y,'w':w,'h':h,'color':c})
def lines_fit(content,width,size,bold):
 font='DVB' if bold else 'DV';out=[]
 for para in content.split('\n'):
  if not para:out.append('');continue
  current=''
  for word in para.split():
   candidate=(current+' '+word).strip()
   if pdfmetrics.stringWidth(candidate,font,size)>width*72-3 and current:out.append(current);current=word
   else:current=candidate
  out.append(current)
 return out

def text(s,x,y,w,h,t,size=18,color=NAVY,bold=False):
 original=size
 lines=lines_fit(t,w,size,bold)
 needed=len(lines)*size*1.22/72
 while needed>h+.01 and size>10:
  size-=.5;lines=lines_fit(t,w,size,bold);needed=len(lines)*size*1.22/72
 if needed>h+.01:raise ValueError(f'Overflow slide {len(slides)} text {t[:50]}')
 if size<16 and original>=16:print('Reduced text',len(slides),size,t[:40])
 s['elements'].append({'type':'text','x':x,'y':y,'w':w,'h':h,'text':'\n'.join(lines),'size':size,'color':color,'bold':bold})
def card(s,x,y,w,h,label,body,accent=TEAL,size=18):
 rect(s,x,y,w,h,LIGHT);rect(s,x,y,.045,h,accent)
 text(s,x+.18,y+.16,w-.36,.45,label,17,accent,True)
 text(s,x+.18,y+.73,w-.36,h-.83,body,size,NAVY)
def bullets(s,x,y,w,items,size=19,gap=.76):
 for i,item in enumerate(items):
  text(s,x,y+i*gap,w,gap-.08,'• '+item,size)

s=new('Stability without concealment',source='Proposal + executed preliminary pilot | Core question inspired by Krishna et al. and groupShapley [S03, S04]',notes='Opening, about 60 seconds. My project asks whether grouping correlated inputs makes model explanations more reliable, or only makes them look more consistent. Imagine six monthly bill amounts: models can distribute importance among these related measurements differently after retraining. I want to distinguish a useful group summary from a summary that hides information. I narrowed the initial COMET proposal because temporal recourse, conformal margins and feature selection would require different evidence. The scope is now feasible in eight weeks on a laptop. The work completed is the proposal audit, literature review and a reproducible logistic-regression pilot; the full experiment is not complete. Name and roll-number fields are editable. Do not call this a finished paper or a new Shapley algorithm.')
text(s,.6,1.65,11.9,.75,'When does grouping correlated inputs improve explanation stability without hiding useful information?',25,NAVY,True)
card(s,.6,2.8,3.9,2.35,'THE PROBLEM','Similar predictions can produce different feature rankings after retraining.',size=20)
card(s,4.7,2.8,3.9,2.35,'THE TRAP','Shorter lists and random grouping can also appear more stable.',accent=GOLD,size=20)
card(s,8.8,2.8,3.9,2.35,'OUR TEST','Compare matched groups; measure stability plus fidelity and cancellation.',size=20)
text(s,.65,5.55,12,.5,'Credit benchmark • interpretable baseline • controlled comparisons • eight weeks',19,TEAL,True)
text(s,.65,6.23,12,.42,'[Student name]  |  [Roll number]  |  B.Tech IT, DTU  |  Prof. Rahul Gupta',17,MUTED)

s=new('What the literature establishes — and what we will test',source='Sources: SHAP [S01]; disagreement [S03]; groupShapley [S04]; ITNAC 2025 [N03]; agreement 2026 [N04]',notes='About 75 seconds. SHAP explains a specified model output using a defined coalition-value game. Krishna and colleagues show that different explanation methods can disagree. groupShapley already explains groups, so grouping is not my invention. I also found a 2025 ITNAC paper on feature grouping for robust SHAP and a 2026 paper on agreement and faithfulness-weighted consensus. Their full methods need to be compared before a novelty claim. Our proposed question is whether correlation grouping beats random partitions with the same group sizes, while preserving fidelity and internal information. We will use synthetic cases with a known explanation game and public credit data. The contribution is an empirical evaluation protocol and findings, not an assertion that combining existing tools creates a new theory. Primary links: https://arxiv.org/abs/2106.12228 ; https://arxiv.org/abs/2202.01602 ; https://doi.org/10.1109/ITNAC66378.2025.11302604 .')
rows=[('SHAP + dependence','Credit depends on model, background and explanation game.','Lundberg & Lee; Aas et al.'),('Disagreement + aggregation','Agreement is a research problem; fusion already exists.','Krishna et al.; Bhatt et al.'),('Grouped explanations','groupShapley and recent robustness work are close prior art.','Jullum et al.; Hendro et al. (2025)')]
for i,(a,b,c) in enumerate(rows):
 y=1.65+i*1.12;rect(s,.6,y,12.1,1.0,LIGHT)
 text(s,.8,y+.15,3.0,.65,a,18,TEAL,True)
 text(s,4,y+.12,8.4,.38,b,17)
 text(s,4,y+.57,8.4,.25,c,12,MUTED)
rect(s,.6,5.24,12.1,1.22,NAVY)
text(s,.85,5.42,11.6,.75,'Proposed gap: test stability beyond matched random grouping, together with information loss and known-ground-truth fidelity.',21,WHITE,True)

s=new('Methodology and controlled evaluation',source='Core method: train-only correlation grouping; coalition sums; size-matched controls. No temporal or causal recourse claim.',notes='About 90 seconds. We split data before fitting preprocessing or clusters. We retrain several models and explain the same held-out individuals. Each explanation is a vector of contributions in a stated output scale. We summarize either individual features or groups formed from training correlation. We compare the resulting stability to random groups with exactly the same number and sizes of groups. That controls a major shorter-list artefact. Next, we examine cancellation, synthetic exact group values and perturbation fidelity so a stable summary cannot simply hide opposing components. The initial model is logistic regression, where interventional logit contributions have an exact closed form. A small tree model and TreeSHAP are planned. The primary partition stays fixed across retraining; reclustering uncertainty is a separate analysis. We keep data, explainer and retraining randomness distinct. Removing model features, causal recourse and drift guarantees are outside the minimum project.')
steps=[('01','Split + preprocess','Fit only on training data'),('02','Retrain + explain','Same held-out people'),('03','Group + controls','Match group sizes exactly'),('04','Evaluate','Stability AND retained detail')]
for i,(n,a,b) in enumerate(steps):
 x=.6+i*3.08;rect(s,x,1.75,2.88,1.65,LIGHT)
 text(s,x+.15,1.9,2.55,.4,n,20,TEAL,True)
 text(s,x+.15,2.38,2.6,.35,a,16,NAVY,True)
 text(s,x+.15,2.87,2.55,.35,b,13,MUTED)
card(s,.6,3.8,5.9,2.45,'MEASURE','RBO stability • matched-random difference\nCancellation • synthetic attribution error\nAP / ROC-AUC / Brier • runtime',size=18)
card(s,6.72,3.8,5.95,2.45,'CONTROL','No grouping • same-size random groups\nKnown additive + interaction cases\nFixed model/output/background per comparison',size=18)
text(s,.65,6.47,12,.28,'Minimum scope: synthetic data + UCI credit; logistic + one tree family. HELOC is conditional on verified access.',13,MUTED)

s=new('Pilot results and eight-week validation plan',source='Measured run: code/results/pilot.json | 10 bootstrap fits; 45 dependent pairs; 30 random partitions | Preliminary only',notes='About 90 seconds. The official UCI dataset contains 30,000 rows. Our pilot used 24,000 training and 6,000 test rows; 500 held-out rows were explained for ten bootstrap logistic models. Correlation grouping produced eleven groups from nineteen inputs. Its ranking RBO was .924, compared with .863 for random partitions with the same sizes: an absolute difference of .061. This is not a 6.1 percent improvement in correctness, and it is not a significance claim. Predictive metrics were AP .489, ROC-AUC .713 and Brier .148. Nine mathematical/data-contract checks pass; a repeated run reproduces deterministic outputs. The main limitation is one split and one simple model, including numeric handling of repayment-status codes. The next steps are synthetic ground truth, fidelity/cancellation analysis, a tree model and additional splits/data. I would like approval for this narrowed scope and feedback on whether an empirical evaluation paper meets project requirements. Target a future ICAIF call or a suitable ECML PKDD workshop, with final venue dependent on results.')
text(s,.65,1.65,6.1,.43,'11 groups; matched group sizes',19,NAVY,True)
for y,label,val,color in [(2.26,'Correlation groups',.9240,TEAL),(3.08,'Matched random groups',.8631,MUTED)]:
 text(s,.65,y,5.7,.3,label,15,MUTED)
 rect(s,.65,y+.38,5.3,.24,LIGHT);rect(s,.65,y+.38,5.3*val,.24,color)
 text(s,6.05,y+.21,.7,.5,f'{val:.3f}',18,color,True)
text(s,.65,4.05,6,.6,'Controlled RBO difference: +0.061',23,TEAL,True)
text(s,.65,4.83,6,1.0,'30,000 rows • 24k train / 6k test\nAP 0.489 • AUC 0.713 • Brier 0.148\n9 checks pass; deterministic replay matches',16)
card(s,7.05,1.65,5.55,3.95,'EIGHT-WEEK VALIDATION','W1–2  Data + synthetic oracle\nW3–4  Tree SHAP + matched controls\nW5–6  Fidelity + replication\nW7–8  Paper + advisor review',size=19)
rect(s,.6,6.0,12.0,.63,'FFF2DD')
text(s,.8,6.14,11.6,.36,'Limit: one split, logistic model only. No significance, causal recourse or publication claim.',15,GOLD,True)

s=new('The concepts you need before the experiment','BACKUP 01 • FIRST PRINCIPLES',notes='Use this if asked to define terms. A feature is an input column, a label is the observed outcome, and a model is a learned mapping. Local attribution concerns one row; global attribution aggregates rows. Stability, fidelity and usefulness are distinct. The full study guide begins before these terms and includes examples.')
card(s,.6,1.7,5.9,2.1,'MODEL','Inputs → learned function → default probability. Testing asks whether predictions generalize.',size=19)
card(s,6.7,1.7,5.9,2.1,'ATTRIBUTION','A contribution to a stated model output, relative to a specified background.',size=19)
card(s,.6,4.05,5.9,2.1,'STABILITY','Does the explanation change when the model is retrained on resampled data?',size=19)
card(s,6.7,4.05,5.9,2.1,'FIDELITY','Does the explanation reflect the specified model/game? Stability alone is insufficient.',size=19)

s=new('The “corrected” manual still needs correction','BACKUP 02 • DOCUMENT AUDIT',source='Detailed 36-point audit with printed PDF page numbers is in the Research Dossier.',notes='The manual contains many sensible corrections but is not fully reliable. Lowering rho merges more groups, not fewer. The proposed quantile code can select the wrong order statistic. Alpha=0 is not a no-margin setting. Mean probability is not sigmoid of mean logit. An additive model does not automatically imply the needed additive conditional attribution game. PFI need not always understate correlated features. Historical repayment data do not support feasible causal recourse. These are reasons for a narrower design, not reasons to discard all explainability research.')
bullets(s,.7,1.8,11.8,[
 'Lower correlation cutoff → more merging, not more groups.',
 'Coalition sums are not generally Owen / exact group values.',
 'Mean logit transformed by sigmoid is not mean probability.',
 'Margin code has an order-statistic error; alpha direction is reversed.',
 'Static credit data cannot validate temporal, actionable recourse.',
 'Agreement, monotonicity and lower VIF do not prove explanation quality.'
],size=20,gap=.76)

s=new('A linear attribution you can calculate by hand','BACKUP 03 • EXACT BASELINE',source='Illustrative arithmetic, not a measured applicant. Pilot uses the same identity on standardized logit inputs.',notes='For score=3*x1+2*x2, reference means 1 and 2 give base 7. At x=(2,5), the feature differences contribute 3 and 6. They sum to the nine-point score difference, and base plus contributions equals sixteen. In logistic regression this is the logit score, not probability. The actual pilot calculates the background mean after each model’s fitted preprocessor and checks this identity for all 500 audit cases. This verifies implementation of a stated attribution, not that the model is causally correct.')
text(s,.75,1.8,11.8,.8,'Score = 3x₁ + 2x₂     Reference = (1, 2)     Input = (2, 5)',25,NAVY,True)
card(s,.7,3.0,3.8,2.1,'BASE SCORE','3×1 + 2×2 = 7',25 if False else TEAL,size=23)
card(s,4.75,3.0,3.8,2.1,'CONTRIBUTIONS','3×(2−1) = 3\n2×(5−2) = 6',size=22)
card(s,8.8,3.0,3.8,2.1,'RECONSTRUCTION','7 + 3 + 6 = 16',size=23)
text(s,.8,5.6,11.8,.65,'Always state the explained output scale. The pilot explains log-odds, not percentage points.',21,TEAL,True)

s=new('Stable summaries can conceal different things','BACKUP 04 • INFORMATION LOSS',source='Coalition sum versus grouped game: groupShapley [S04]. Worked numbers are illustrative.',notes='Two different issues are shown. Cancellation: +5 and minus 4.8 net to .2, while the unsigned total is 9.8. Net contribution is legitimate but can hide strong opposing components. Game semantics: in a three-player unanimity game every feature gets one third; summing A and B gives two thirds. Treat AB as one group and C as another, and each group gets one half. So the order of grouping and Shapley computation matters. We retain within-group drill-down and use exact small-game values as an oracle.')
card(s,.6,1.8,5.9,3.6,'CANCELLATION','Feature credits: +5 and −4.8\nNet group contribution: 0.2\nUnsigned internal mass: 9.8\nA stable net can hide strong components.',size=20)
card(s,6.7,1.8,5.9,3.6,'GROUPING CHANGES THE GAME','Three-player unanimity game:\nordinary A + B = 2/3\nmerged {A,B} group = 1/2\nCoalition sum ≠ exact group value.',size=20)
text(s,.65,5.8,12,.65,'Retain group summaries AND feature drill-down; test interactions with a small exact oracle.',22,TEAL,True)

s=new('Pilot provenance and what its numbers mean','BACKUP 05 • ACTUAL EXECUTED WORK',source='Official UCI dataset 350, CC BY 4.0 | https://doi.org/10.24432/C55S3H | SHA256 and versions in pilot.json',notes='The official dataset was downloaded during this task and hashed. Ten bootstrap models share the same held-out test set. Forty-five pairwise comparisons share models and are not independent experimental units. Feature RBO .776 has nineteen items and group RBO .924 has eleven, so that raw difference is not the main controlled comparison. Random groups preserve the exact group sizes and average .863. This pilot does not implement the full proposed train/development/test protocol; it uses a prespecified 80/20 split without tuning. Numeric repayment-status encoding remains a known limitation.')
card(s,.6,1.75,5.9,4.35,'DATA + MODEL','30,000 observed customer records\n19 numeric inputs after exclusions\n24,000 train / 6,000 test\n500 held-out audit rows\n10 bootstrap logistic models\nExact interventional logit attribution',size=18)
card(s,6.7,1.75,5.9,4.35,'MEASUREMENTS','19-feature RBO: 0.776\n11-group RBO: 0.924\nSize-matched random RBO: 0.863\nControlled difference: +0.061\nMax reconstruction error: 2.22e−15\nNo inferential p-value claimed',size=18)
text(s,.65,6.35,12,.3,'Single split + simple model + fixed partition. Fidelity, second dataset and full benchmark remain unrun.',14,GOLD,True)

s=new('Full experiment: isolate one source of change at a time','BACKUP 06 • RESEARCH DESIGN',notes='For the final study, save stratified train/development/test IDs. Train-only transforms and clusters prevent leakage. The primary experiment freezes the correlation partition while retraining models. A separate reclustering analysis measures group discovery uncertainty. A separate explainer-repeat experiment measures sampling variance. Add a synthetic known-game oracle and one tree family. HELOC is conditional on obtaining a verified release and dictionary. The study can succeed as a negative-result evaluation if gains vanish under controls; do not keep expanding modules to force a positive result.')
bullets(s,.7,1.75,11.9,[
 'Data: controlled synthetic cases + UCI; verified HELOC if available.',
 'Models: logistic regression first; one small tree ensemble next.',
 'Split: 60% train / 20% development / 20% locked test.',
 'Separate: model resampling, explainer sampling and reclustering.',
 'Controls: singleton, random same-size, domain and giant groups.',
 'Small exact group games verify additive and interaction cases.'
],size=20,gap=.77)

s=new('Evaluation must measure more than consistency','BACKUP 07 • METRICS',notes='RBO is the primary ranking stability measure. The main effect subtracts the mean of matched random partitions, but this is still a descriptive effect until replicated. Cancellation captures net-versus-unsigned attribution compression. Synthetic oracle error compares against a defined exact game, not assumed causal truth. AP, AUC and Brier validate the predictor; they do not validate explanations. Borrower confidence intervals with fixed models are conditional intervals, not a universal measure of model uncertainty. Report all denominators and failures.')
for i,(a,b) in enumerate([
 ('Stability','RBO at p=.85; sensitivity at .7/.95; fixed audit individuals.'),
 ('Controlled gain','Correlation-group stability minus size-matched random-group stability.'),
 ('Retained information','Cancellation, drill-down, exact synthetic error and masker sensitivity.'),
 ('Model + uncertainty','AP / AUC / Brier; paired effects; honest conditional uncertainty.')]):
 y=1.7+i*1.18;rect(s,.65,y,12,1.02,LIGHT)
 text(s,.85,y+.15,3.15,.6,a,19,TEAL,True);text(s,4.3,y+.15,8.05,.65,b,18)

s=new('What would falsify the project hypothesis?','BACKUP 08 • NEGATIVE RESULTS',notes='A useful project must be able to fail its hypothesis. If matched random groups are just as stable, there is little evidence that correlation-specific grouping is the cause. If group summaries hide major opposing contributions or fail known-game fidelity checks, stability is insufficient. If benefit exists only in one model/split, generality is weak. If closest prior work already includes the proposed controls, novelty must be revised. These cases can support a careful replication or limitation study, but they do not guarantee publication.')
card(s,.6,1.8,5.9,3.8,'EVIDENCE AGAINST','Random groups perform equally well.\nInformation loss offsets stability.\nResults disappear on new splits.\nClosest prior art already answers the question.',accent=GOLD,size=20)
card(s,6.7,1.8,5.9,3.8,'HONEST RESPONSE','Report the failure conditions.\nKeep all runs and controls.\nRevise claims and scope.\nDo not invent expected results or tune on test.',size=20)
text(s,.65,6.0,12,.5,'Research quality comes from a valid answer, not from forcing a positive outcome.',22,TEAL,True)

s=new('The code demonstration: two minutes, no hand-waving','BACKUP 09 • IMPLEMENTATION',notes='Show demo.py, test_demo.py and results/pilot.json. Explain the data exclusions and split, the linear attribution formula, and the matched random partitions. Run the tests; nine should pass. Open the saved result and identify model/data counts and limitations. The quantile helper is only an audit correction and is not used by the main pipeline. The starter includes a downloader and Colab notebook, but no personal Colab account run has been performed. Be candid about AI assistance and your responsibility for validation.')
card(s,.6,1.8,5.9,3.9,'SHOW THE FILES','demo.py — measured baseline\ntest_demo.py — 9 checks\ndownload_data.py — official input\nresults/pilot.json — traceable evidence\nColab_Quickstart.ipynb — CPU route',size=18)
card(s,6.7,1.8,5.9,3.9,'EXPLAIN THE CHECKS','Disjoint training/test IDs\nComplete-linkage condition\nHand-computed RBO and linear values\nExact logit reconstruction\nIdentical deterministic rerun',size=18)
text(s,.65,6.08,12,.42,'Not implemented yet: tree SHAP, LIME, full synthetic oracle, recourse, second-data validation.',17,GOLD,True)

s=new('Eight weeks, with a deliverable gate each fortnight','BACKUP 10 • EXECUTION PLAN',notes='Week 1 learn and reproduce baseline; week 2 finish data cards and synthetic oracle. Weeks 3 and 4 add the tree explanation and controlled stability experiment. Weeks 5 and 6 assess fidelity, cancellation and replication. Weeks 7 and 8 write and independently replay the work with advisor feedback. Optional features are dropped if core gates fail. Use bounded Cursor tasks; do not ask it to generate a whole paper and system in one pass. Free Colab CPU is sufficient for the scoped experiments.')
for i,(a,b) in enumerate([
 ('WEEKS 1–2','Understand baseline; audit data; verify exact synthetic games.'),
 ('WEEKS 3–4','Add tree SHAP; run matched-group stability and negative controls.'),
 ('WEEKS 5–6','Measure information loss; replicate; freeze main results.'),
 ('WEEKS 7–8','Write, review, replay and prepare a venue-specific submission.')]):
 y=1.8+i*1.13;rect(s,.65,y,12,1.0,LIGHT)
 text(s,.85,y+.2,2.65,.4,a,18,TEAL,True);text(s,3.75,y+.16,8.5,.68,b,19)

s=new('Publication plan: target fit, not a promise','BACKUP 11 • VENUES',source='Verified series pages: icaif2026.org • ecmlpkdd.org • pakdd.org | No upcoming deadline commitment',notes='ICAIF is the best topical fit among these targets for credit explainability, but the 2026 main submissions are closed and workshop submissions closed October 12. Its 2026 event is November 14 to 17 in Milan. We would target a future call, plausibly 2027, after verifying the actual CFP. An appropriate ECML PKDD workshop may fit an evaluation/reproducibility study, but its archival publication policy must satisfy DTU requirements. PAKDD is a stretch if the methodological contribution becomes strong. EJOR, DSS, ESWA and KBS listed in the old manual are journals, not conferences. There is no guarantee of acceptance and no fee/payment/submission has been made.')
card(s,.6,1.85,3.9,3.55,'PRIMARY FIT','Future ACM ICAIF call\nFinancial ML + explainability\n2026 submissions closed\nVerify next CFP with advisor',size=19)
card(s,4.7,1.85,3.9,3.55,'EVALUATION OPTION','ECML PKDD XAI / trustworthy-ML workshop\nCheck actual workshop and archival policy',size=19)
card(s,8.8,1.85,3.9,3.55,'STRETCH','PAKDD research track\nRequires stronger methodological evidence\nNo deadline verified',size=19)
text(s,.65,5.85,12,.65,'Ask the supervisor: Is the narrowed empirical study suitable, and what publication format does DTU require?',21,TEAL,True)

s=new('Selected references and the novelty checkpoint','BACKUP 12 • REFERENCES',source='Full linked bibliography, verification levels and recent-work audit are in the Research Dossier.',notes='Read the closest sources rather than memorizing citation counts. SHAP defines the attribution framework; Krishna documents disagreement; groupShapley is the central grouping baseline; Chen and colleagues explain model-versus-data semantics. The 2025 feature-grouping robustness paper and 2026 consensus/agreement paper require closest-work comparison before claiming novelty. Do not claim to have reviewed their entire full text if only metadata or abstract was accessible. The research dossier includes more than thirty linked primary sources and distinguishes unresolved references.')
refs=[
 'Lundberg & Lee (2017). A Unified Approach to Interpreting Model Predictions. arXiv:1705.07874.',
 'Krishna et al. The Disagreement Problem in Explainable Machine Learning. arXiv:2202.01602.',
 'Jullum, Redelmeier & Aas. groupShapley. arXiv:2106.12228.',
 'Chen et al. (2020). True to the Model or True to the Data? arXiv:2006.16234.',
 'Hendro, Mei & Welch (2025). Feature-Grouping for Robust SHAP. DOI: 10.1109/ITNAC66378.2025.11302604.',
 'Quantifying overall agreement… (2026). DOI: 10.1007/s12065-026-01145-y.'
]
bullets(s,.7,1.75,11.9,refs,size=17,gap=.78)

# Render both editable PPTX and matching PDF from the same scene data.
def output(n,stem):
 prs=Presentation();prs.slide_width=Inches(W);prs.slide_height=Inches(H)
 pdf=canvas.Canvas(str(B/(stem+'.pdf')),pagesize=(W*72,H*72))
 pdf.setTitle(stem.replace('_',' '))
 for s in slides[:n]:
  sl=prs.slides.add_slide(prs.slide_layouts[6])
  for e in s['elements']:
   x,y,w,h=[e[a] for a in ['x','y','w','h']]
   if e['type']=='rect':
    shape=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(e['color']);shape.line.fill.background()
    pdf.setFillColor('#'+e['color']);pdf.rect(x*72,(H-y-h)*72,w*72,h*72,stroke=0,fill=1)
   else:
    box=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=box.text_frame
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.word_wrap=False
    for i,line in enumerate(e['text'].split('\n')):
     p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line
     p.font.name='DejaVu Sans';p.font.size=Pt(e['size']);p.font.bold=e['bold'];p.font.color.rgb=RGBColor.from_string(e['color'])
     p.space_after=Pt(0);p.space_before=Pt(0);p.line_spacing=Pt(e['size']*1.22)
     pdf.setFont('DVB' if e['bold'] else 'DV',e['size']);pdf.setFillColor('#'+e['color'])
     pdf.drawString(x*72,(H-y)*72-e['size']*.97-i*e['size']*1.22,line)
  sl.notes_slide.notes_text_frame.text=s['notes']
  pdf.showPage()
 prs.save(B/(stem+'.pptx'));pdf.save()
output(4,'04_Presentation_4_Slides')
output(len(slides),'05_Extended_Deck_16_Slides')
with (B/'06_Presentation_Speaking_Notes.md').open('w') as f:
 f.write('# Presentation speaking notes\n\nThe first four slides form a roughly five-minute proposal presentation. The remaining slides are optional backups. Replace editable placeholders. All pilot numbers are preliminary.\n')
 for i,s in enumerate(slides,1):f.write(f'\n## Slide {i}: {s["title"]}\n\n{s["notes"]}\n')
(B/'build/slide_scenes.json').write_text(json.dumps(slides,indent=2))
print('Created',len(slides),'slides; 4-slide and 16-slide PPTX/PDF; notes embedded.')
