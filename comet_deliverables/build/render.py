from pathlib import Path
import re, html, textwrap, subprocess, json
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Preformatted, KeepTogether
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from docx import Document
from docx.shared import Pt, Inches as DI, RGBColor as DC
B=Path(__file__).resolve().parent.parent
pdfmetrics.registerFont(TTFont('DV','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
pdfmetrics.registerFont(TTFont('DV-Bold','/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'))
pdfmetrics.registerFont(TTFont('DV-Mono','/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf'))
pdfmetrics.registerFontFamily('DV',normal='DV',bold='DV-Bold',italic='DV',boldItalic='DV-Bold')
NAVY='#11283F'; TEAL='#087F8C'; MUTED='#53667A'
styles={
 'p':ParagraphStyle('p',fontName='DV',fontSize=9.5,leading=14,spaceAfter=7,textColor=colors.HexColor(NAVY),splitLongWords=True),
 'h1':ParagraphStyle('h1',fontName='DV-Bold',fontSize=25,leading=30,spaceBefore=5,spaceAfter=16,textColor=colors.HexColor(NAVY),keepWithNext=True),
 'h2':ParagraphStyle('h2',fontName='DV-Bold',fontSize=15,leading=19,spaceBefore=15,spaceAfter=9,textColor=colors.HexColor(TEAL),keepWithNext=True),
 'h3':ParagraphStyle('h3',fontName='DV-Bold',fontSize=11.5,leading=16,spaceBefore=11,spaceAfter=7,textColor=colors.HexColor(NAVY),keepWithNext=True),
 'code':ParagraphStyle('code',fontName='DV-Mono',fontSize=7.7,leading=10.5,spaceAfter=10,backColor=colors.HexColor('#EDF2F6'),borderPadding=7),
}
def markup(s):
 s=html.escape(s)
 s=re.sub(r'\*\*(.*?)\*\*',r'<b>\1</b>',s)
 s=re.sub(r'`([^`]+)`',r'<font name="DV-Mono">\1</font>',s)
 return s

def make_pdf(path):
 lines=path.read_text().splitlines();story=[];i=0
 while i<len(lines):
  line=lines[i]
  if not line.strip():i+=1;continue
  if line.startswith('```'):
   i+=1;code=[]
   while i<len(lines) and not lines[i].startswith('```'):
    code.extend(textwrap.wrap(lines[i],width=95,replace_whitespace=False,drop_whitespace=False) or ['']);i+=1
   story.append(Preformatted('\n'.join(code),styles['code']));i+=1;continue
  m=re.match(r'^(#{1,3}) (.*)',line)
  if m:story.append(Paragraph(markup(m[2]),styles['h'+str(len(m[1]))]));i+=1;continue
  para=[line];i+=1
  while i<len(lines) and lines[i].strip() and not re.match(r'^(#|```|- |\d+\. )',lines[i]):
   para.append(lines[i]);i+=1
  text=' '.join(para)
  if text.startswith('- '):text='• '+text[2:]
  story.append(Paragraph(markup(text),styles['p']))
 def page(c,d):
  c.saveState();c.setFont('DV',8);c.setFillColor(colors.HexColor(MUTED))
  c.drawString(42, A4[1]-28,'COMET-XAI LITE  |  '+path.stem.replace('_',' '))
  c.drawString(42,24,'DTU • Proposal and preliminary pilot • 27 October 2026')
  c.drawRightString(A4[0]-42,24,str(d.page));c.restoreState()
 doc=SimpleDocTemplate(str(path.with_suffix('.pdf')),pagesize=A4,rightMargin=42,leftMargin=42,topMargin=48,bottomMargin=42,title=path.stem,author='Prepared for DTU final-year project')
 doc.build(story,onFirstPage=page,onLaterPages=page)

def make_docx(path):
 out=path.with_suffix('.docx')
 subprocess.run(['pandoc',str(path),'-o',str(out),'--standalone'],check=True)
 d=Document(out)
 normal=d.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10.5)
 normal.paragraph_format.space_after=Pt(7)
 for name,size in [('Title',27),('Heading 1',22),('Heading 2',16),('Heading 3',12)]:
  
  if name not in d.styles: continue
  d.styles[name].font.name='Calibri';d.styles[name].font.size=Pt(size);d.styles[name].font.color.rgb=DC.from_string('11283F' if name!='Heading 2' else '087F8C')
 for sec in d.sections:
  sec.top_margin=DI(.7);sec.bottom_margin=DI(.7);sec.left_margin=DI(.75);sec.right_margin=DI(.75)
  sec.header.paragraphs[0].text='COMET-XAI Lite | DTU final-year research project'
  sec.footer.paragraphs[0].text='Proposal + preliminary pilot • 27 October 2026'
 d.save(out)
for p in sorted(B.glob('0[123]_*.md')):
 make_pdf(p);make_docx(p);print(p.stem,'PDF + DOCX complete')
